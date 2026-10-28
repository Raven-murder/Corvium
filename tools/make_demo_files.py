from pathlib import Path
import textwrap


def main():
    from docx import Document
    from pptx import Presentation
    from PIL import Image, ImageDraw, ImageFont

    folder = Path(__file__).resolve().parents[1] / "demo_data"
    folder.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("A transaction protects related changes", 0)
    document.add_paragraph(
        "A club transfers 40 credits from its equipment budget to its travel budget. "
        "The debit and credit belong in one transaction. If either update fails, "
        "rollback prevents the accounts from recording only half of the transfer."
    )
    document.add_heading("Check before commit", level=1)
    document.add_paragraph(
        "Confirm that both budgets exist and that equipment has enough credits. "
        "Protect the balance check against concurrent spending. Commit only after "
        "both updates succeed. A successful transaction preserves the combined balance."
    )
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Outcome"
    table.rows[0].cells[1].text = "Required action"
    for left, right in (("Both updates succeed", "Commit"), ("One update fails", "Roll back")):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = left, right
    document.save(folder / "transactions_sample.docx")

    presentation = Presentation()
    slides = [
        ("One logical operation", "Transfer 40 credits between two budgets.\nTreat the debit and credit as one transaction.",
         "The transfer should preserve the budgets' combined balance."),
        ("Commit or roll back", "Commit after both updates succeed.\nRoll back if a required update fails.",
         "Atomicity prevents a partially recorded transfer."),
        ("Concurrent spending", "Protect the balance check and debit together.\nKeep transactions short and handle conflicts.",
         "A second request must not spend credits already reserved by the first request."),
    ]
    for title, body, notes in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
    presentation.save(folder / "transactions_slides.pptx")

    with Image.new("RGB", (1200, 500), "white") as image:
        try:
            font = ImageFont.load_default(size=30)
        except TypeError:
            font = ImageFont.load_default()
        text = "DATABASE TRANSACTIONS\n\n" + textwrap.fill(
            "A transfer debits one budget and credits another. Commit when both "
            "updates succeed. Roll back when either update fails. The combined "
            "balance should remain unchanged. Protect balance checks against "
            "concurrent spending.", width=55,
        )
        ImageDraw.Draw(image).multiline_text((35, 35), text, fill="black", font=font, spacing=16)
        image.save(folder / "transactions_text.png")


if __name__ == "__main__":
    main()
