from io import BytesIO
import importlib.util
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import numpy as np
from PIL import Image
from pypdf import PdfWriter

with patch("dotenv.load_dotenv"):
    from src import config, db, demo, extractors, ingest, ocr, workspaces


class ExtractorTests(unittest.TestCase):
    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx is not installed")
    def test_docx_paragraphs_tables_and_sections(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Introduction")
        document.add_heading("Atomicity", level=1)
        document.add_paragraph("Commit both updates together.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Failure"
        table.cell(0, 1).text = "Rollback"
        document.add_heading("Durability", level=2)
        document.add_paragraph("Committed changes survive a restart.")
        stream = BytesIO()
        document.save(stream)
        stream.seek(0)
        sections = extractors.extract_docx(stream)
        self.assertEqual([part["section"] for part in sections], [None, "Atomicity", "Durability"])
        self.assertIn("Failure\tRollback", sections[1]["text"])
        self.assertLess(sections[1]["text"].index("Commit"), sections[1]["text"].index("Rollback"))
        self.assertIn("survive a restart", sections[2]["text"])
        self.assertTrue(all(part["page"] == 1 for part in sections))
        self.assertEqual(extractors.extract_docx(stream.getvalue()), sections)

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx is not installed")
    def test_pptx_slide_pages_and_notes(self):
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Atomicity"
        slide.placeholders[1].text = "Commit both updates."
        slide.notes_slide.notes_text_frame.text = "Roll back if one update fails."
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Isolation"
        slide.placeholders[1].text = "Coordinate concurrent requests."
        stream = BytesIO()
        presentation.save(stream)
        stream.seek(0)
        pages = extractors.extract_pptx(stream)
        self.assertEqual([part["page"] for part in pages], [1, 2])
        self.assertEqual([part["section"] for part in pages], ["Atomicity", "Isolation"])
        self.assertIn("Commit both updates.", pages[0]["text"])
        self.assertIn("Roll back if one update fails.", pages[0]["text"])
        self.assertNotIn("Roll back", pages[1]["text"])
        self.assertEqual(extractors.extract_pptx(stream.getvalue()), pages)

    def test_html_visible_text_and_headings(self):
        contents = b"""<html><head><title>Hidden title</title><style>hidden CSS</style></head>
        <body><p>Introduction &amp; scope.</p><h1>Atomic <em>updates</em></h1>
        <p>Commit <b>both</b> changes.</p><script>hidden script</script>
        <div hidden><p>hidden attribute</p></div><div style="display: none">hidden style</div>
        <h2>Recovery</h2><p>Roll back on failure.</p></body></html>"""
        sections = extractors.extract_html(BytesIO(contents).getvalue())
        self.assertEqual([part["section"] for part in sections], [None, "Atomic updates", "Recovery"])
        combined = "\n".join(part["text"] for part in sections)
        self.assertIn("Introduction & scope.", combined)
        self.assertIn("Commit both changes.", sections[1]["text"])
        self.assertNotIn("hidden", combined.lower())
        self.assertIn("Roll back on failure.", sections[2]["text"])

    def test_html_heading_breaks_and_empty_attributes(self):
        sections = extractors.extract_html(b'<h1>Atomic<br>updates</h1><p style aria-hidden>Keep both changes.</p>')
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0]["section"], "Atomic updates")
        self.assertIn("Keep both changes.", sections[0]["text"])

    def test_markdown_nearest_heading_and_code_fence(self):
        contents = BytesIO(b"Introduction\n# Atomicity\nCommit together.\n```python\n# a code comment\n```\nRecovery\n--------\nRollback.\n")
        sections = extractors.extract_text(contents.getvalue())
        self.assertEqual([part["section"] for part in sections], [None, "Atomicity", "Recovery"])
        self.assertIn("# a code comment", sections[1]["text"])
        self.assertIn("Rollback.", sections[2]["text"])
        chunks = ingest.chunk_text(sections[1]["text"], section=sections[1]["section"], chunk_size=20, overlap=3)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk["section"] == "Atomicity" for chunk in chunks))

    def test_text_encoding_fallback(self):
        self.assertEqual(extractors.extract_text("# Café\nRésumé".encode("utf-8"))[0]["section"], "Café")
        self.assertIn("Résumé", extractors.extract_text("# Café\nRésumé".encode("latin-1"))[0]["text"])


class IngestionTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.folder = Path(temp.name)
        storage = workspaces.storage_context(self.folder)
        storage.__enter__()
        self.addCleanup(storage.__exit__, None, None, None)
        embed = patch.object(ingest, "embed_texts", side_effect=lambda texts: np.ones((len(texts), 4096), dtype=np.float32))
        self.embed = embed.start()
        self.addCleanup(embed.stop)
        db.init_db()

    def document(self, name):
        with db.database_session() as connection:
            return dict(connection.execute("SELECT * FROM documents WHERE name = ?", (name,)).fetchone())

    def test_stores_original_kind_and_chunk_sections(self):
        path = self.folder / "notes.md"
        path.write_bytes(b"# Atomicity\nCommit both updates.\n# Recovery\nRoll back together.")
        self.assertEqual(ingest.ingest_file(path), 2)
        document = self.document(path.name)
        self.assertEqual(db.get_file_bytes(db.current_workspace_id(), document["id"]), path.read_bytes())
        self.assertEqual(document["kind"], "text")
        self.assertEqual(document["status"], "indexed")
        with db.database_session() as connection:
            self.assertEqual([row[0] for row in connection.execute("SELECT section FROM chunks ORDER BY id")], ["Atomicity", "Recovery"])

    def test_no_text_and_ocr_not_configured_are_recorded(self):
        path = self.folder / "scan.png"
        path.write_bytes(b"original image bytes")
        with patch.object(ingest, "get_vision_model", return_value=""), patch.object(ingest, "ocr_image") as recognize:
            self.assertEqual(ingest.ingest_file(path), 0)
            recognize.assert_not_called()
        document = self.document(path.name)
        self.assertEqual(document["kind"], "image")
        self.assertIn("no text found", document["status"])
        self.assertIn("OCR not configured", document["status"])
        self.assertEqual(db.get_file_bytes(db.current_workspace_id(), document["id"]), path.read_bytes())
        self.embed.assert_not_called()
        empty = self.folder / "empty.txt"
        empty.write_bytes(b"")
        self.assertEqual(ingest.ingest_file(empty), 0)
        self.assertEqual(self.document(empty.name)["status"], "no text found")

    def test_image_ocr_text_is_indexed(self):
        path = self.folder / "scan.webp"
        path.write_bytes(b"webp fixture")
        with patch.object(ingest, "get_vision_model", return_value="vision-model"), patch.object(ingest, "ocr_image", return_value="Commit both updates") as recognize:
            self.assertEqual(ingest.ingest_file(path), 1)
            recognize.assert_called_once_with(b"webp fixture")
        self.assertEqual(self.document(path.name)["kind"], "image")

    def test_batch_continues_after_oversize_or_bad_file(self):
        large = self.folder / "large.txt"
        with large.open("wb") as handle:
            handle.truncate(ingest.MAX_FILE_BYTES + 1)
        broken = self.folder / "broken.pdf"
        broken.write_bytes(b"Not a PDF")
        valid = self.folder / "notes.log"
        valid.write_bytes("Résumé about transactions".encode("latin-1"))
        with patch.object(ingest, "PdfReader", side_effect=ValueError("Invalid PDF")):
            results = ingest.ingest_files([large, broken, valid])
        self.assertIn("25 MB", results[0]["error"])
        self.assertIsNotNone(results[1]["error"])
        self.assertIsNone(results[2]["error"])
        self.assertEqual(results[2]["count"], 1)
        self.assertEqual(self.document(valid.name)["kind"], "text")
        with db.database_session() as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0], 1)

    def test_failed_reindex_preserves_previous_original_and_chunks(self):
        path = self.folder / "notes.md"
        path.write_bytes(b"First version")
        ingest.ingest_file(path)
        document = self.document(path.name)
        path.write_bytes(b"Changed version")
        with patch.object(ingest, "put_file", side_effect=db.StorageError()):
            with self.assertRaises(db.StorageError):
                ingest.ingest_file(path)
        self.assertEqual(self.document(path.name), document)
        self.assertEqual(db.get_file_bytes(db.current_workspace_id(), document["id"]), b"First version")
        with db.database_session() as connection:
            self.assertEqual(connection.execute("SELECT text FROM chunks").fetchone()[0], "First version")

    def test_pdf_ocr_only_for_short_pages(self):
        pages = [SimpleNamespace(extract_text=lambda: "A sufficiently long paragraph containing database notes."),
                 SimpleNamespace(extract_text=lambda: "Short")]
        with patch.object(ingest, "PdfReader", return_value=SimpleNamespace(pages=pages)), patch.object(ingest, "get_vision_model", return_value="vision-model"), patch.object(ingest, "_render_pdf_page", return_value=b"rendered PNG") as render, patch.object(ingest, "ocr_image", return_value="Transcribed database notes") as recognize:
            result = ingest._extract_contents(b"pdf fixture", ".pdf")
        self.assertEqual([part["page"] for part in result], [1, 2])
        self.assertEqual(result[1]["text"], "Transcribed database notes")
        render.assert_called_once_with(b"pdf fixture", 1)
        recognize.assert_called_once_with(b"rendered PNG")
        with patch.object(ingest, "PdfReader", return_value=SimpleNamespace(pages=pages)), patch.object(ingest, "get_vision_model", return_value=""), patch.object(ingest, "_render_pdf_page") as render:
            result = ingest._extract_contents(b"pdf fixture", ".pdf")
        render.assert_not_called()
        self.assertEqual(result[1]["text"], "Short")
        self.assertEqual(result[1]["warnings"], ["OCR not configured"])

    @unittest.skipUnless(importlib.util.find_spec("pypdfium2"), "pypdfium2 is not installed")
    def test_pdfium_renders_a_page_to_png(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=100)
        output = BytesIO()
        writer.write(output)
        rendered = ingest._render_pdf_page(output.getvalue(), 0)
        with Image.open(BytesIO(rendered)) as image:
            self.assertEqual(image.format, "PNG")
            self.assertEqual(image.size, (400, 200))

    def test_demo_discovers_new_formats(self):
        folder = self.folder / "demo_data"
        folder.mkdir()
        for name in ("a.md", "b.docx", "c.pptx", "d.png", "ignore.json"):
            (folder / name).write_bytes(b"fixture")
        with patch.object(demo, "PROJECT_ROOT", self.folder), patch.object(demo, "save_memories"), patch.object(demo, "ingest_files", return_value=[]) as batch:
            demo._seed_demo_files()
        self.assertEqual([path.suffix for path in batch.call_args.args[0]], [".md", ".docx", ".pptx", ".png"])


class OCRTests(unittest.TestCase):
    def test_ocr_skip_and_shared_model_request(self):
        with patch.object(ocr, "get_vision_model", return_value=""), patch.object(ocr, "get_client") as client:
            self.assertEqual(ocr.ocr_image(b"unused"), "")
            client.assert_not_called()
        output = BytesIO()
        with Image.new("RGB", (20, 20), "white") as image:
            image.save(output, format="PNG")
        response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="OCR TEST 123"))])
        with patch.object(ocr, "get_vision_model", return_value="configured-model"), patch.object(ocr, "get_client") as client:
            client.return_value.chat.completions.create.return_value = response
            self.assertEqual(ocr.ocr_image(output.getvalue()), "OCR TEST 123")
            request = client.return_value.chat.completions.create.call_args.kwargs
        self.assertEqual(request["model"], "configured-model")
        self.assertIn("Transcribe all visible text exactly", request["messages"][0]["content"][0]["text"])
        self.assertTrue(request["messages"][0]["content"][1]["image_url"]["url"].startswith("data:image/png;base64,"))


if __name__ == "__main__":
    unittest.main()
