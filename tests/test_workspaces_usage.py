import ast
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import importlib.util
import os
from pathlib import Path
import sqlite3
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

import numpy as np

with patch("dotenv.load_dotenv"):
    from src import actions, chats, config, db, demo, ingest, library, memory, qa, sources, usage, workspaces

A = "a" * 32
B = "b" * 32


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.data = self.root / "data"
        demo_folder = self.root / "demo_data"
        demo_folder.mkdir()
        (demo_folder / "sample.md").write_text("# Transactions\nA transaction commits related changes together.", encoding="utf-8")
        for target, name, value in ((config, "DATA_DIR", self.data), (demo, "PROJECT_ROOT", self.root)):
            mocked = patch.object(target, name, value)
            mocked.start()
            self.addCleanup(mocked.stop)
        embedding = patch.object(ingest, "embed_texts", side_effect=lambda texts: np.ones((len(texts), 4096), dtype=np.float32))
        self.embed = embedding.start()
        self.addCleanup(embedding.stop)
        memory_embedding = patch.object(memory, "embed_texts", side_effect=lambda texts: np.eye(4096, dtype=np.float32)[:len(texts)])
        self.memory_embed = memory_embedding.start()
        self.addCleanup(memory_embedding.stop)
        vision = patch.object(ingest, "get_vision_model", return_value="")
        vision.start()
        self.addCleanup(vision.stop)

    def context(self, workspace_id):
        storage = workspaces.workspace_storage(workspace_id)
        workspaces.ensure_workspace(storage)
        return workspaces.storage_context(storage.root, workspace_id)

    def test_database_access_without_workspace_fails_closed(self):
        with self.assertRaises(workspaces.WorkspaceError):
            db.connect()
        self.assertFalse((self.data / "app.db").exists())

    def test_invalid_ids_are_rejected_before_paths_are_built(self):
        for value in (None, "", "../demo_template", "A" * 32, "g" * 32, A + "/../" + B, "a" * 31, [A]):
            with self.subTest(value=value):
                with self.assertRaises(workspaces.WorkspaceError):
                    workspaces.workspace_storage(value)
                self.assertRegex(workspaces.resolve_workspace_id(value), r"^[0-9a-f]{32}$")
        self.assertEqual(workspaces.resolve_workspace_id(A), A)
        self.assertFalse(self.data.exists())
        with self.assertRaises(workspaces.WorkspaceError):
            with workspaces.storage_context(self.root, A):
                pass

    def test_binding_preserves_id_on_refresh_and_clears_cross_workspace_state(self):
        query, state = {}, {"quiz": "Old result", "selected_chat_id": 12}
        workspace_id, unchanged = workspaces.bind_workspace(query, state)
        self.assertFalse(unchanged)
        self.assertEqual(query["w"], workspace_id)
        self.assertEqual(state, {"workspace_id": workspace_id})
        self.assertEqual(workspaces.bind_workspace(query, state), (workspace_id, True))
        state["destructive_confirmation"] = {"kind": "document"}
        query["w"] = B
        self.assertEqual(workspaces.bind_workspace(query, state), (B, False))
        self.assertEqual(state, {"workspace_id": B})

    def test_callback_does_not_reuse_another_workspaces_widget_arguments(self):
        tree = ast.parse((Path(__file__).resolve().parents[1] / "app.py").read_text(encoding="utf-8"))
        node = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == "run_callback")
        state = {"workspace_id": A, "destructive_confirmation": {"item": 1}}
        ui = SimpleNamespace(query_params={"w": B}, session_state=state)
        namespace = {"st": ui, "bind_workspace": workspaces.bind_workspace, "safe_error": str}
        exec(compile(ast.Module(body=[node], type_ignores=[]), "<workspace-callback>", "exec"), namespace)
        action = Mock()
        namespace["run_callback"](action, 1)
        action.assert_not_called()
        self.assertEqual(state, {"workspace_id": B})
        namespace["run_callback"](action, 2)
        action.assert_called_once_with(state, 2)

    def test_concurrent_workspace_contexts_do_not_leak(self):
        workspaces.ensure_template()
        def visit(workspace_id):
            with self.context(workspace_id):
                chat_id = chats.create_chat(workspace_id)
                chats.add_message(chat_id, "user", workspace_id)
                return workspaces.get_storage().workspace_id, [row["content"] for row in chats.list_messages(chat_id)]
        with ThreadPoolExecutor(max_workers=2) as pool:
            result = list(pool.map(visit, (A, B)))
        self.assertEqual(result, [(A, [A]), (B, [B])])

    def test_workspaces_isolate_documents_memories_chats_and_originals(self):
        with self.context(A):
            storage_a = workspaces.get_storage()
            path = storage_a.uploads_dir / "private-a.md"
            path.parent.mkdir()
            path.write_text("Private A content.", encoding="utf-8")
            ingest.ingest_file(path)
            memory.add_memory("Private A memory")
            chat_a = chats.create_chat("Private A chat")
            chats.add_message(chat_a, "user", "Private A message")
            library.create_folder("Private A folder")
            original = next(row for row in library.list_documents() if row["name"] == path.name)
            original_path = workspaces.resolve_original_path(original["stored_path"])
        with self.context(B):
            self.assertNotIn("private-a.md", [row["name"] for row in library.list_documents()])
            self.assertNotIn("Private A memory", [row["text"] for row in memory.list_memories()])
            self.assertNotIn("Private A chat", [row["title"] for row in chats.list_chats()])
            self.assertNotIn("Private A folder", [row["name"] for row in library.list_library()["folders"]])
            chat_b = chats.create_chat("Private B chat")
            chats.add_message(chat_b, "user", "Private B message")
            self.assertEqual([row["content"] for row in chats.list_messages(chat_b)], ["Private B message"])
            self.assertIsNone(sources.original_download({"stored_path": str(original_path), "name": "private-a.md"}))
            with self.assertRaises(workspaces.WorkspaceError):
                ingest.ingest_file(path)
        with self.context(A):
            self.assertEqual([row["content"] for row in chats.list_messages(chat_a)], ["Private A message"])
            self.assertNotIn("Private B chat", [row["title"] for row in chats.list_chats()])
        self.assertNotEqual(workspaces.workspace_storage(A).db_path, workspaces.workspace_storage(B).db_path)
        self.assertFalse((self.data / "app.db").exists())

    def test_template_is_built_once_then_copied_without_embedding_calls(self):
        with patch.object(demo, "_seed_demo_files", wraps=demo._seed_demo_files) as seed:
            template = workspaces.ensure_template()
            embedded = self.embed.call_count
            memories_embedded = self.memory_embed.call_count
            with self.context(A):
                document = library.list_documents()[0]
                self.assertFalse(Path(document["stored_path"]).is_absolute())
                self.assertTrue(workspaces.resolve_original_path(document["stored_path"]).is_file())
            with self.context(B):
                self.assertEqual(len(library.list_documents()), 1)
                self.assertEqual(len(memory.list_memories()), 2)
            self.assertEqual(self.embed.call_count, embedded)
            self.assertEqual(self.memory_embed.call_count, memories_embedded)
            seed.assert_called_once()
            self.assertTrue((template / ".ready").is_file())

    @unittest.skipUnless(importlib.util.find_spec("docx") and importlib.util.find_spec("pptx"), "Office parsers are not installed")
    def test_template_includes_office_files_and_skips_unconfigured_image_ocr(self):
        from docx import Document
        from pptx import Presentation
        from PIL import Image

        folder = self.root / "demo_data"
        document = Document()
        document.add_paragraph("Commit both updates.")
        document.save(folder / "sample.docx")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Atomicity"
        presentation.save(folder / "sample.pptx")
        with Image.new("RGB", (40, 40), "white") as image:
            image.save(folder / "sample.png")
        with patch.object(ingest, "ocr_image") as ocr:
            with self.context(A):
                documents = library.list_documents()
                self.assertEqual({row["kind"] for row in documents}, {"text", "docx", "pptx", "image"})
                image = next(row for row in documents if row["kind"] == "image")
                self.assertIn("OCR not configured", image["status"])
                self.assertTrue(all(workspaces.resolve_original_path(row["stored_path"]).is_file() for row in documents))
            ocr.assert_not_called()

    def test_simultaneous_first_visitors_publish_one_complete_template(self):
        with patch.object(demo, "_seed_demo_files", wraps=demo._seed_demo_files) as seed:
            with ThreadPoolExecutor(max_workers=2) as pool:
                list(pool.map(lambda value: workspaces.ensure_workspace(workspaces.workspace_storage(value)), (A, B)))
            seed.assert_called_once()
        for workspace_id in (A, B):
            with self.context(workspace_id):
                self.assertEqual(len(library.list_documents()), 1)
                with db.database_session() as connection:
                    self.assertEqual(connection.execute("PRAGMA integrity_check").fetchone()[0], "ok")

    def test_failed_template_build_is_not_published(self):
        with patch.object(demo, "_seed_demo_files", side_effect=RuntimeError("Build interrupted")):
            with self.assertRaises(RuntimeError):
                workspaces.ensure_template()
        self.assertFalse(workspaces.template_directory().exists())
        self.assertEqual(list(self.data.glob(".demo-template-*")), [])
        workspaces.ensure_template()
        self.assertTrue((workspaces.template_directory() / ".ready").is_file())

    def test_reset_restores_only_current_workspace_without_resetting_usage(self):
        with self.context(A):
            chats.create_chat("Extra chat")
            memory.add_memory("Private A fact")
            with patch.object(config, "WORKSPACE_CHAT_LIMIT", 1):
                usage.reserve_workspace_action("chat")
                before = self.embed.call_count
                demo.reset_demo()
                self.assertEqual(self.embed.call_count, before)
                self.assertEqual(chats.list_chats(), [])
                self.assertNotIn("Private A fact", [row["text"] for row in memory.list_memories()])
                with self.assertRaises(usage.UsageLimitError):
                    usage.reserve_workspace_action("chat")
        with self.context(B):
            chats.create_chat("Keep B")
        with self.context(A):
            demo.reset_demo()
        with self.context(B):
            self.assertEqual(chats.list_chats()[0]["title"], "Keep B")

    def test_cleanup_removes_only_workspaces_older_than_seven_days(self):
        for workspace_id in (A, B):
            workspaces.ensure_workspace(workspaces.workspace_storage(workspace_id))
        now = 2_000_000
        for workspace_id, days in ((A, 8), (B, 6)):
            marker = workspaces.workspace_storage(workspace_id).root / ".last_used"
            os.utime(marker, (now - days * 86400, now - days * 86400))
        unrelated = self.data / "workspaces" / "not-a-workspace"
        unrelated.mkdir()
        workspaces.cleanup_workspaces(now=now)
        self.assertFalse(workspaces.workspace_storage(A).root.exists())
        self.assertTrue(workspaces.workspace_storage(B).root.exists())
        self.assertTrue(workspaces.template_directory().is_dir())
        self.assertTrue(unrelated.exists())
        workspaces.cleanup_workspaces(now=now + 20 * 86400, exclude=(B,))
        self.assertTrue(workspaces.workspace_storage(B).root.exists())

    def test_workspace_categories_and_rolling_windows_are_enforced(self):
        limits = {"WORKSPACE_CHAT_LIMIT": 2, "WORKSPACE_TOOL_LIMIT": 1, "WORKSPACE_UPLOAD_LIMIT": 1}
        with patch.multiple(config, **limits):
            with self.context(A):
                for kind, allowed in (("chat", 2), ("tool", 1), ("upload", 1)):
                    for _ in range(allowed):
                        usage.reserve_workspace_action(kind, now=1000)
                    with self.assertRaises(usage.UsageLimitError):
                        usage.reserve_workspace_action(kind, now=1001)
                usage.reserve_workspace_action("chat", now=87400)
            with self.context(B):
                usage.reserve_workspace_action("chat", now=87400)
        self.assertTrue((self.data / "usage.db").is_file())

    def test_global_api_limit_is_shared_atomic_and_has_a_private_message(self):
        with patch.object(config, "GLOBAL_API_ACTION_LIMIT", 5):
            def attempt(_):
                try:
                    usage.reserve_api_action(now=1000)
                    return True
                except usage.UsageLimitError as error:
                    message = config.safe_error(error)
                    self.assertNotRegex(message, r"[0-9/\\]")
                    return False
            with ThreadPoolExecutor(max_workers=8) as pool:
                results = list(pool.map(attempt, range(16)))
            self.assertEqual(sum(results), 5)
            usage.reserve_api_action(now=87400)

    def test_sdk_guard_checks_every_call_before_sending(self):
        chat_call, embedding_call, model_call = Mock(), Mock(), Mock()
        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=chat_call)),
                                 embeddings=SimpleNamespace(create=embedding_call), models=SimpleNamespace(list=model_call))
        with patch.object(config, "GLOBAL_API_ACTION_LIMIT", 2), patch.object(config, "get_api_key", return_value="fixture"), patch.object(config, "OpenAI", return_value=client):
            guarded = config.get_client()
            guarded.embeddings.create(input=["text"])
            guarded.chat.completions.create(messages=[])
            with self.assertRaises(usage.UsageLimitError):
                guarded.models.list()
        chat_call.assert_called_once()
        embedding_call.assert_called_once()
        model_call.assert_not_called()

    def test_limits_block_chat_tool_and_upload_before_api_work(self):
        with self.context(A):
            path = workspaces.get_storage().root / "upload.md"
            path.write_text("Text to index.", encoding="utf-8")
            with patch.multiple(config, WORKSPACE_CHAT_LIMIT=0, WORKSPACE_TOOL_LIMIT=0, WORKSPACE_UPLOAD_LIMIT=0), patch.object(qa, "_answer_question") as answer, patch.object(actions, "_complete") as complete:
                for function in (lambda: qa.answer_question("Question"), lambda: actions.generate_study_plan(), lambda: actions.generate_quiz(), lambda: ingest.ingest_file(path)):
                    with self.assertRaises(usage.UsageLimitError):
                        function()
                answer.assert_not_called()
                complete.assert_not_called()


if __name__ == "__main__":
    unittest.main()
