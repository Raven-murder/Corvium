from html.parser import HTMLParser
from io import BytesIO
import re


def decode_text(contents):
    try:
        return contents.decode("utf-8-sig")
    except UnicodeDecodeError:
        return contents.decode("latin-1")


def _sections(blocks, page=1):
    result = []
    for section, text in blocks:
        text = text.strip()
        if not text:
            continue
        if result and result[-1]["section"] == section:
            result[-1]["text"] += "\n" + text
        else:
            result.append({"page": page, "section": section, "text": text})
    return result


def extract_text(contents):
    text = contents if isinstance(contents, str) else decode_text(contents)
    lines = text.splitlines()
    blocks, section, fence = [], None, None
    index = 0
    while index < len(lines):
        line = lines[index]
        marker = re.match(r"^\s{0,3}(`{3,}|~{3,})", line)
        if marker:
            if fence is None:
                fence = marker.group(1)[0]
            elif marker.group(1)[0] == fence:
                fence = None
            blocks.append((section, line))
            index += 1
            continue
        heading = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$", line) if fence is None else None
        underline = (
            fence is None and line.strip() and index + 1 < len(lines)
            and re.fullmatch(r"\s{0,3}(?:=+|-+)\s*", lines[index + 1])
        )
        if heading:
            section = heading.group(1)
            blocks.append((section, section))
        elif underline:
            section = line.strip()
            blocks.append((section, section))
            index += 1
        else:
            blocks.append((section, line))
        index += 1
    return _sections(blocks)


def extract_docx(source):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(source) if isinstance(source, bytes) else source)
    blocks, section = [], None
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            style = paragraph.style.name if paragraph.style is not None else ""
            if text and re.match(r"^(?:heading\s*\d+|title)$", style, re.I):
                section = text
            blocks.append((section, text))
        elif tag == "tbl":
            table = Table(element, document)
            for row in table.rows:
                blocks.append((section, "\t".join(cell.text for cell in row.cells)))
    return _sections(blocks)


def _shape_text(shapes):
    result = []
    for shape in shapes:
        if getattr(shape, "has_text_frame", False):
            result.append(shape.text_frame.text)
        if getattr(shape, "has_table", False):
            result.extend("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)
        if hasattr(shape, "shapes"):
            result.extend(_shape_text(shape.shapes))
    return result


def extract_pptx(source):
    from pptx import Presentation

    presentation = Presentation(BytesIO(source) if isinstance(source, bytes) else source)
    result = []
    for number, slide in enumerate(presentation.slides, 1):
        parts = _shape_text(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None:
                parts.append(notes.text)
        title = slide.shapes.title
        result.append({
            "page": number, "section": title.text.strip() if title is not None else None,
            "text": "\n".join(part.strip() for part in parts if part.strip()),
        })
    return result


class _VisibleHTML(HTMLParser):
    _VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}
    _BLOCK = {"p", "div", "section", "article", "li", "ul", "ol", "table", "tr", "td", "th", "pre", "blockquote", "br", "hr"}
    _HIDDEN = {"script", "style", "head", "template", "noscript"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack = []
        self.blocks = []
        self.parts = []
        self.section = None
        self.heading = None

    def _flush(self):
        text = " ".join("".join(self.parts).split())
        if text:
            if self.heading:
                self.section = text
            self.blocks.append((self.section, text))
        self.parts = []

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        hidden = (
            bool(self.stack and self.stack[-1][1]) or tag in self._HIDDEN
            or "hidden" in attributes or (attributes.get("aria-hidden") or "").lower() == "true"
            or bool(re.search(r"(?:display\s*:\s*none|visibility\s*:\s*hidden)", (attributes.get("style") or ""), re.I))
        )
        if not hidden:
            if self.heading and tag in self._BLOCK:
                self.parts.append(" ")
            elif tag in self._BLOCK or re.fullmatch(r"h[1-6]", tag):
                self._flush()
            if re.fullmatch(r"h[1-6]", tag):
                self.heading = tag
        if tag not in self._VOID:
            self.stack.append((tag, hidden))

    def handle_endtag(self, tag):
        if not (self.stack and self.stack[-1][1]):
            if tag == self.heading:
                self._flush()
                self.heading = None
            elif tag in self._BLOCK and not self.heading:
                self._flush()
        for index in range(len(self.stack) - 1, -1, -1):
            if self.stack[index][0] == tag:
                del self.stack[index:]
                break

    def handle_data(self, data):
        if not (self.stack and self.stack[-1][1]):
            self.parts.append(data)


def extract_html(contents):
    parser = _VisibleHTML()
    parser.feed(contents if isinstance(contents, str) else decode_text(contents))
    parser.close()
    parser._flush()
    return _sections(parser.blocks)
